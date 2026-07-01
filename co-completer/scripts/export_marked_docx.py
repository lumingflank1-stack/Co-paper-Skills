#!/usr/bin/env python3
"""Export Markdown with red bold revision spans to DOCX.

Recognized addition markup:
  <span style="color:red"><strong>text</strong></span>
  <span style="color:red"><b>text</b></span>
  {{ADD:text}}
  <red>text</red>

The script intentionally implements a compact Markdown subset so it can run
without pandoc. It handles headings, paragraphs, simple lists, blockquotes,
simple pipe tables, fenced code blocks, bold text, and red-bold additions.
"""

from __future__ import annotations

import argparse
import html
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
except ImportError as exc:  # pragma: no cover - environment message
    raise SystemExit(
        "python-docx is required. Try the bundled Python runtime or install python-docx."
    ) from exc


RED = RGBColor(255, 0, 0)

TOKEN_RE = re.compile(
    r"(<span[^>]*color\s*:\s*red[^>]*>\s*<(?:strong|b)>(.*?)</(?:strong|b)>\s*</span>"
    r"|<red>(.*?)</red>"
    r"|\{\{ADD:(.*?)\}\}"
    r"|\*\*(.*?)\*\*"
    r"|<(?:strong|b)>(.*?)</(?:strong|b)>)",
    re.IGNORECASE | re.DOTALL,
)

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
TAG_RE = re.compile(r"<[^>]+>")


def clean_text(text: str) -> str:
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    text = TAG_RE.sub("", text)
    return html.unescape(text)


def add_inline_runs(paragraph, text: str) -> None:
    pos = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(clean_text(text[pos : match.start()]))

        red_text = match.group(2) or match.group(3) or match.group(4)
        bold_text = match.group(5) or match.group(6)

        if red_text is not None:
            run = paragraph.add_run(clean_text(red_text))
            run.bold = True
            run.font.color.rgb = RED
        elif bold_text is not None:
            run = paragraph.add_run(clean_text(bold_text))
            run.bold = True

        pos = match.end()

    if pos < len(text):
        paragraph.add_run(clean_text(text[pos:]))


def split_table_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [cell.strip() for cell in line.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def add_table(document: Document, rows: list[list[str]]) -> None:
    rows = [row for row in rows if row]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(width):
            cell = table.cell(r_idx, c_idx)
            text = row[c_idx] if c_idx < len(row) else ""
            paragraph = cell.paragraphs[0]
            add_inline_runs(paragraph, text)
            if r_idx == 0:
                for run in paragraph.runs:
                    run.bold = True


def set_base_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"


def convert(markdown_path: Path, docx_path: Path, title: str | None = None) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    document = Document()
    set_base_styles(document)

    if title:
        document.add_heading(title, level=0)

    lines = text.splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []

    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = document.add_paragraph(style="Intense Quote")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not stripped:
            idx += 1
            continue

        if "|" in stripped and stripped.startswith("|"):
            table_rows: list[list[str]] = []
            while idx < len(lines):
                candidate = lines[idx].strip()
                if not (candidate.startswith("|") and "|" in candidate):
                    break
                if not is_table_separator(candidate):
                    table_rows.append(split_table_row(candidate))
                idx += 1
            add_table(document, table_rows)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            paragraph = document.add_heading(level=level)
            add_inline_runs(paragraph, heading.group(2))
            idx += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, bullet.group(1))
            idx += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline_runs(paragraph, numbered.group(1))
            idx += 1
            continue

        quote = re.match(r"^>\s?(.*)$", stripped)
        if quote:
            paragraph = document.add_paragraph(style="Intense Quote")
            add_inline_runs(paragraph, quote.group(1))
            idx += 1
            continue

        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, stripped)
        idx += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path, help="Input Markdown file")
    parser.add_argument("docx", type=Path, help="Output DOCX file")
    parser.add_argument("--title", help="Optional document title")
    args = parser.parse_args(argv)

    if not args.markdown.exists():
        parser.error(f"Input Markdown does not exist: {args.markdown}")

    convert(args.markdown, args.docx, args.title)
    print(f"Wrote {args.docx}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
